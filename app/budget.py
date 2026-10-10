"""Reads the budgeting spreadsheet (the source of truth for budget amounts).

The sheet is laid out as small tables, each with a title, a header row
(label, Week, Fortnight, Month, ...) and one row per line. A separate
"Subscriptions" sheet lists services in (name, Price Per Month) column pairs.
The file is re-read only when it changes on disk.
"""

from dataclasses import dataclass, field
from pathlib import Path

from . import config


@dataclass
class Budget:
    sheet: str = ""
    # Table title -> {line name: amount per fortnight}, e.g. "Bank Accounts" -> {"Food": 200.0}
    tables: dict[str, dict[str, float]] = field(default_factory=dict)
    # [{"name", "category", "monthly"}] from the Subscriptions sheet
    subscriptions: list[dict] = field(default_factory=list)
    error: str | None = None

    def table(self, title: str) -> dict[str, float]:
        for name, rows in self.tables.items():
            if name.lower() == title.lower():
                return rows
        return {}


_cache: tuple[float, Budget] | None = None


def load() -> Budget:
    global _cache
    if not config.BUDGET_FILE:
        return Budget(error="Set BUDGET_FILE in .env to the path of your budgeting spreadsheet.")
    path = Path(config.BUDGET_FILE).expanduser()
    try:
        mtime = path.stat().st_mtime
    except OSError:
        return Budget(error=f"Can't find the budget spreadsheet at {path}.")
    if _cache and _cache[0] == mtime:
        return _cache[1]
    try:
        budget = _read(path)
    except Exception as exc:  # a half-synced or open file shouldn't break the page
        return Budget(error=f"Couldn't read the budget spreadsheet: {exc}")
    _cache = (mtime, budget)
    return budget


def _read(path: Path) -> Budget:
    import openpyxl

    # data_only gives the values Excel last calculated, not the formulas.
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = {ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    wb.close()

    if config.BUDGET_SHEET:
        if config.BUDGET_SHEET not in sheets:
            raise ValueError(f'there\'s no sheet called "{config.BUDGET_SHEET}" (BUDGET_SHEET in .env)')
        name = config.BUDGET_SHEET
    else:
        name = next((n for n, rows in sheets.items() if "Bank Accounts" in _tables(rows)), next(iter(sheets)))

    subs = next((rows for n, rows in sheets.items() if n.lower().startswith("subscription")), [])
    return Budget(sheet=name, tables=_tables(sheets[name]), subscriptions=_subscriptions(subs))


def _text(value) -> str:
    return value.strip() if isinstance(value, str) else ""


def _tables(rows: list[list]) -> dict[str, dict[str, float]]:
    """Find every header row with a "Fortnight" column and read the table under it."""
    tables: dict[str, dict[str, float]] = {}
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            if _text(cell) != "Fortnight" or c < 2 or _text(row[c - 1]) != "Week":
                continue
            label_col = c - 2
            title = _title_above(rows, r, label_col, c + 5)
            lines: dict[str, float] = {}
            for below in rows[r + 1:]:
                label = _text(below[label_col]) if label_col < len(below) else ""
                value = below[c] if c < len(below) else None
                if not label:
                    break
                if label.lower().startswith("total") or label == "Account":
                    continue
                if isinstance(value, (int, float)):
                    lines[label] = round(float(value), 2)
            tables[title] = lines
    return tables


def _title_above(rows: list[list], r: int, first_col: int, last_col: int) -> str:
    for above in (rows[r - 1] if r >= 1 else [], rows[r - 2] if r >= 2 else []):
        for cell in above[first_col:last_col + 1]:
            if _text(cell) and _text(cell) not in ("Account",):
                return _text(cell)
    # The Bank Accounts table repeats its "Account" label on the header row itself.
    return _text(rows[r][first_col]) or f"Table {r + 1}"


def _subscriptions(rows: list[list]) -> list[dict]:
    """(name, Price Per Month) column pairs; the header of the name column is the category."""
    if not rows:
        return []
    header = rows[0]
    items = []
    for c in range(len(header) - 1):
        if "price" not in _text(header[c + 1]).lower():
            continue
        category = _text(header[c]) or "Other"
        for row in rows[1:]:
            label = _text(row[c]) if c < len(row) else ""
            price = row[c + 1] if c + 1 < len(row) else None
            if label.lower().startswith("total"):
                break
            if label and isinstance(price, (int, float)):
                items.append({"name": label, "category": category, "monthly": round(float(price), 2)})
    return items
